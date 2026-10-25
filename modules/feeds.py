"""
feeds.py

Single access point for the seven CDC surveillance feeds and the curated
workbook. A feed is looked up in this order:

  1. <data dir>/feeds/<sheet>.parquet   written by scripts/refresh_data.py
     (locally, in Colab, or by the weekly GitHub workflow);
  2. the sheet of the legacy master workbook "Research Data.xlsx";
  3. the latest weekly snapshot on the repository's data branch, downloaded
     once and cached on local disk (web deployment); the manifest is
     rechecked every REMOTE_TTL seconds, so a new week appears without a
     restart.

A missing feed returns an empty frame and is reported by data_status();
no page fails because a feed is absent.
"""

import json
import os
import time
from pathlib import Path
from typing import Dict, Optional

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent

FEED_SHEETS = [
    "Weekly Hospital Respiratory Dat",
    "NSSP Emergency Department Visit",
    "Percent of Tests Positive for V",
    "CDC Wastewater Data for SARS-Co",
    "CDC Epidemic Trends and Rt",
    "RESP-NET Rates and Clinical Dat",
    "NNDSS Weekly Data",
]
REMOTE_TTL = 3 * 3600


def _secret(name: str, default: str = "") -> str:
    v = os.environ.get(name)
    if v:
        return v
    try:
        import streamlit as st
        return str(st.secrets.get(name, default))
    except Exception:
        return default


def data_dir() -> Path:
    return Path(os.environ.get("SSDM_DATA_DIR", ROOT / "data"))


def cache_dir() -> Path:
    env = os.environ.get("SSDM_CACHE_DIR")
    d = Path(env) if env else (Path("/content/ssdm_cache") if Path("/content").exists()
                               else ROOT / ".cache")
    d.mkdir(parents=True, exist_ok=True)
    return d


def workbook_path() -> Optional[Path]:
    for p in [data_dir() / "Research Data.xlsx", ROOT / "data" / "Research Data.xlsx",
              ROOT / "Research Data.xlsx"]:
        if p.exists():
            return p
    return None


def curated_path() -> Path:
    """Curated sheets (baseline evidence, registry, dictionary, citations,
    derivations): the curated seed in the repository is the master copy; a
    legacy workbook is used only when the seed is absent."""
    seed = ROOT / "data" / "curated_seed.xlsx"
    return seed if seed.exists() else (workbook_path() or seed)


def _fname(sheet: str) -> str:
    return sheet.replace("/", "_") + ".parquet"


def _normalize(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip().lower().replace(" ", "_").replace("-", "_") for c in df.columns]
    return df


# ----------------------------------------------------------------------
# remote data branch
# ----------------------------------------------------------------------

def _remote_cfg():
    repo = _secret("SSDM_DATA_REPO", "Shayan9574/ssdm-dashboard")
    branch = _secret("SSDM_DATA_BRANCH", "data")
    token = _secret("GITHUB_TOKEN", "")
    return repo, branch, token


def _remote_get(path: str) -> Optional[bytes]:
    import requests
    repo, branch, token = _remote_cfg()
    url = f"https://raw.githubusercontent.com/{repo}/{branch}/{path}"
    headers = {"Authorization": f"token {token}"} if token else {}
    try:
        r = requests.get(url, headers=headers, timeout=60)
        if r.status_code == 200:
            return r.content
    except Exception:
        return None
    return None


def _remote_dir() -> Path:
    d = cache_dir() / "remote"
    d.mkdir(parents=True, exist_ok=True)
    return d


def sync_remote(force: bool = False) -> Optional[dict]:
    """Downloads the data branch manifest and any changed feed file."""
    d = _remote_dir()
    stamp = d / ".checked"
    man_file = d / "manifest.json"
    if (not force and stamp.exists() and man_file.exists()
            and time.time() - stamp.stat().st_mtime < REMOTE_TTL):
        return json.loads(man_file.read_text())
    raw = _remote_get("manifest.json")
    if raw is None:
        return json.loads(man_file.read_text()) if man_file.exists() else None
    man = json.loads(raw)
    old = json.loads(man_file.read_text()) if man_file.exists() else {}
    for name, meta in man.get("files", {}).items():
        local = d / name
        if local.exists() and old.get("files", {}).get(name, {}).get("sha256") == meta.get("sha256"):
            continue
        blob = _remote_get(f"feeds/{name}") if not name.startswith("derived/") else _remote_get(name)
        if blob is not None:
            local.parent.mkdir(parents=True, exist_ok=True)
            local.write_bytes(blob)
    man_file.write_text(json.dumps(man, indent=2))
    stamp.write_text(str(time.time()))
    return man


# ----------------------------------------------------------------------
# feed access
# ----------------------------------------------------------------------

_MEMO: Dict[str, tuple] = {}
_SHEETS: Dict[str, tuple] = {}


def _sheetnames(wb: Path) -> list:
    key = f"{wb}|{wb.stat().st_mtime_ns}"
    if wb.as_posix() not in _SHEETS or _SHEETS[wb.as_posix()][0] != key:
        try:
            from openpyxl import load_workbook
            names = load_workbook(wb, read_only=True).sheetnames
        except Exception:
            names = []
        _SHEETS[wb.as_posix()] = (key, names)
    return _SHEETS[wb.as_posix()][1]


def _source_for(sheet: str):
    p = data_dir() / "feeds" / _fname(sheet)
    if p.exists():
        return "local parquet", p
    wb = workbook_path()
    if wb is not None and sheet in _sheetnames(wb):
        return "workbook", wb
    if _secret("SSDM_REMOTE_DATA", "1") != "0":
        sync_remote()
        rp = _remote_dir() / _fname(sheet)
        if rp.exists():
            return "data branch", rp
    return "missing", None


def load_feed(sheet: str) -> pd.DataFrame:
    kind, path = _source_for(sheet)
    if path is None:
        return pd.DataFrame()
    stt = path.stat()
    key = f"{sheet}|{path}|{stt.st_mtime_ns}|{stt.st_size}"
    if sheet in _MEMO and _MEMO[sheet][0] == key:
        return _MEMO[sheet][1]
    if kind == "workbook":
        from modules.data_cache import read_excel_sheet
        df = read_excel_sheet(path, sheet_name=sheet, header=1)
    else:
        df = pd.read_parquet(path)
    df = _normalize(df)
    _MEMO[sheet] = (key, df)
    return df


def feed_signature() -> str:
    """Changes whenever any feed file changes; part of every cache key."""
    parts = []
    for s in FEED_SHEETS:
        kind, path = _source_for(s)
        st_ = path.stat() if path else None
        parts.append(f"{s}:{kind}:{st_.st_mtime_ns if st_ else 0}:{st_.st_size if st_ else 0}")
    cp = curated_path()
    parts.append(f"curated:{cp.stat().st_mtime_ns if cp.exists() else 0}")
    return "|".join(parts)


DATE_COLS = {
    "Weekly Hospital Respiratory Dat": "weekendingdate",
    "NSSP Emergency Department Visit": "week_end",
    "Percent of Tests Positive for V": "week_end",
    "CDC Wastewater Data for SARS-Co": "sample_collect_date",
    "CDC Epidemic Trends and Rt": "date",
    "RESP-NET Rates and Clinical Dat": "date",
}


def data_status() -> pd.DataFrame:
    rows = []
    for s in FEED_SHEETS:
        kind, path = _source_for(s)
        df = load_feed(s) if path is not None else pd.DataFrame()
        last = None
        if not df.empty:
            col = DATE_COLS.get(s)
            if col and col in df.columns:
                last = pd.to_datetime(df[col], errors="coerce").max()
            elif {"year", "week"} <= set(df.columns):
                y = pd.to_numeric(df["year"], errors="coerce")
                w = pd.to_numeric(df["week"], errors="coerce")
                m = y.max()
                last = f"{int(m)} week {int(w[y == m].max())}" if m == m else None
        rows.append({"Feed": s, "Source": kind, "Records": len(df),
                     "Latest observation": (str(last)[:10] if not isinstance(last, str) else last) if last is not None else "none"})
    return pd.DataFrame(rows)


def update_log() -> Optional[pd.DataFrame]:
    """The weekly workflow's record of what changed since the previous week."""
    for p in [data_dir() / "derived" / "update_log.csv", _remote_dir() / "derived" / "update_log.csv"]:
        if p.exists():
            try:
                return pd.read_csv(p)
            except Exception:
                return pd.DataFrame()
    return None
